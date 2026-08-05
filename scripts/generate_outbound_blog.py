#!/usr/bin/env python3
"""Generate the Word source for the Azure outbound-connectivity article."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent.parent
POSTS_DIR = ROOT / "posts"
OUTPUT = POSTS_DIR / "disabling-default-outbound-connectivity-in-azure.docx"
ASSET_DIR = ROOT / "assets" / "outbound-connectivity"

NAVY = "17365D"
BLUE = "0078D4"
TEAL = "008272"
PALE_BLUE = "EAF4FB"
PALE_TEAL = "E7F5F2"
GRAY = "5B6573"
LIGHT_GRAY = "E5E7EB"
WHITE = "FFFFFF"


def paint(value: str) -> str:
    return f"#{value}"


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Segoe UI"
    run.font.size = Pt(27)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(subtitle)
    run.bold = True
    run.font.name = "Segoe UI"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)


def add_callout(document: Document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(label.upper() + "\n")
    label_run.bold = True
    label_run.font.color.rgb = RGBColor.from_string(BLUE)
    text_run = paragraph.add_run(text)
    text_run.font.color.rgb = RGBColor.from_string(NAVY)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    properties = paragraph._p.get_or_add_pPr()
    style = properties.find(qn("w:pStyle"))
    if style is None:
        style = OxmlElement("w:pStyle")
        properties.insert(0, style)
    style.set(qn("w:val"), f"Heading{level + 1}")
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Segoe UI"
    run.font.size = Pt(18 if level == 1 else 14)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.add_run("• ").bold = True
        paragraph.add_run(item)


def add_labeled_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(label + ": ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph.add_run(text)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = None
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                set_cell_shading(cells[index], "F7F8FA")
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cells[index].text = value
            if widths:
                cells[index].width = Inches(widths[index])
    document.add_paragraph()


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = (
        [Path("C:/Windows/Fonts/seguisb.ttf"), Path("C:/Windows/Fonts/segoeui.ttf")]
        if bold
        else [Path("C:/Windows/Fonts/segoeui.ttf"), Path("C:/Windows/Fonts/arial.ttf")]
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def rounded_box(draw: ImageDraw.ImageDraw, bounds, fill: str, outline: str, title: str, detail: str) -> None:
    draw.rounded_rectangle(bounds, radius=20, fill=paint(fill), outline=paint(outline), width=4)
    left, top, right, _ = bounds
    draw.text(((left + right) / 2, top + 30), title, font=font(30, True), fill=paint(NAVY), anchor="mm")
    draw.multiline_text(((left + right) / 2, top + 85), detail, font=font(22), fill=paint(GRAY), anchor="ma", align="center", spacing=8)


def draw_arrow(draw: ImageDraw.ImageDraw, start, end, color: str, label: str) -> None:
    draw.line((start, end), fill=paint(color), width=8)
    x, y = end
    draw.polygon([(x, y), (x - 24, y - 14), (x - 24, y + 14)], fill=paint(color))
    midpoint = ((start[0] + end[0]) / 2, start[1] - 30)
    draw.text(midpoint, label, font=font(20, True), fill=paint(color), anchor="mm")


def create_snat_diagram() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "packet-journey-snat.png"
    image = Image.new("RGB", (1800, 760), paint(WHITE))
    draw = ImageDraw.Draw(image)
    draw.text((900, 52), "One outbound HTTPS connection, two identities", font=font(40, True), fill=paint(NAVY), anchor="mm")

    rounded_box(draw, (80, 190, 510, 500), PALE_BLUE, BLUE, "Workload VM", "Private source\n10.20.1.4:51542")
    rounded_box(draw, (685, 190, 1115, 500), PALE_TEAL, TEAL, "Explicit egress", "SNAT state\n10.20.1.4:51542\n→ 203.0.113.10:62001")
    rounded_box(draw, (1290, 190, 1720, 500), "F7F8FA", GRAY, "Public service", "Destination\n198.51.100.25:443")

    draw_arrow(draw, (510, 330), (685, 330), BLUE, "Private packet")
    draw_arrow(draw, (1115, 330), (1290, 330), TEAL, "Translated packet")
    draw.line((1290, 430, 1115, 430), fill=paint(GRAY), width=5)
    draw.polygon([(1115, 430), (1139, 416), (1139, 444)], fill=paint(GRAY))
    draw.line((685, 430, 510, 430), fill=paint(GRAY), width=5)
    draw.polygon([(510, 430), (534, 416), (534, 444)], fill=paint(GRAY))
    draw.text((900, 615), "Return traffic is admitted only because it matches the stored outbound flow state.", font=font(26, True), fill=paint(NAVY), anchor="mm")
    draw.text((900, 665), "The application keeps its private address; the internet sees the explicit egress public address.", font=font(23), fill=paint(GRAY), anchor="mm")
    image.save(path, quality=95)
    return path


def create_horizontal_flow(filename: str, title: str, nodes: list[tuple[str, str, str]], footer: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / filename
    image = Image.new("RGB", (1800, 700), paint(WHITE))
    draw = ImageDraw.Draw(image)
    draw.text((900, 52), title, font=font(40, True), fill=paint(NAVY), anchor="mm")
    count = len(nodes)
    gap = 55
    box_width = int((1640 - gap * (count - 1)) / count)
    left = 80
    for index, (node_title, detail, color) in enumerate(nodes):
        right = left + box_width
        rounded_box(draw, (left, 190, right, 480), "F7F8FA", color, node_title, detail)
        if index < count - 1:
            draw_arrow(draw, (right, 335), (right + gap, 335), color, "")
        left = right + gap
    draw.text((900, 595), footer, font=font(25, True), fill=paint(NAVY), anchor="mm")
    image.save(path, quality=95)
    return path


def create_options_diagram() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "explicit-outbound-options.png"
    image = Image.new("RGB", (1800, 1050), paint(WHITE))
    draw = ImageDraw.Draw(image)
    draw.text((900, 54), "Five explicit ways to reach a destination", font=font(40, True), fill=paint(NAVY), anchor="mm")
    rounded_box(draw, (635, 150, 1165, 380), PALE_BLUE, BLUE, "Private workload", "The source stays private.\nThe architecture chooses the path.")
    options = [
        ((80, 520, 560, 780), "NAT Gateway", "Subnet-wide dynamic SNAT", TEAL),
        ((660, 520, 1140, 780), "Load Balancer", "Backend-pool outbound rules", BLUE),
        ((1240, 520, 1720, 780), "Firewall / NVA", "Central inspection through UDR", "C239B3"),
        ((370, 830, 850, 1010), "Instance public IP", "A narrow 1:1 exception", "D83B01"),
        ((950, 830, 1430, 1010), "Private Link", "Avoid public egress for PaaS", "107C10"),
    ]
    for bounds, title, detail, color in options:
        rounded_box(draw, bounds, "F7F8FA", color, title, detail)
        target_x = (bounds[0] + bounds[2]) // 2
        target_y = bounds[1]
        draw.line((900, 380, target_x, target_y), fill=paint(color), width=5)
    image.save(path, quality=95)
    return path


def create_snat_capacity_diagram() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "snat-capacity-models.png"
    image = Image.new("RGB", (1800, 850), paint(WHITE))
    draw = ImageDraw.Draw(image)
    draw.text((900, 52), "Why port allocation changes the scaling conversation", font=font(40, True), fill=paint(NAVY), anchor="mm")
    draw.rounded_rectangle((70, 145, 865, 750), radius=24, fill=paint(PALE_TEAL), outline=paint(TEAL), width=4)
    draw.rounded_rectangle((935, 145, 1730, 750), radius=24, fill=paint(PALE_BLUE), outline=paint(BLUE), width=4)
    draw.text((468, 195), "NAT Gateway", font=font(32, True), fill=paint(TEAL), anchor="mm")
    draw.text((1332, 195), "Load Balancer outbound rules", font=font(32, True), fill=paint(BLUE), anchor="mm")
    for index, y in enumerate((300, 420, 540)):
        draw.rounded_rectangle((135, y, 330, y + 70), radius=12, fill=paint(WHITE), outline=paint(TEAL), width=3)
        draw.text((232, y + 35), f"VM {index + 1}", font=font(22, True), fill=paint(NAVY), anchor="mm")
        draw.line((330, y + 35, 650, 420), fill=paint(TEAL), width=5)
    draw.rounded_rectangle((650, 335, 800, 505), radius=16, fill=paint(TEAL), outline=paint(TEAL), width=3)
    draw.multiline_text((725, 420), "Shared\ndynamic\npool", font=font(22, True), fill=paint(WHITE), anchor="mm", align="center")
    draw.text((468, 665), "Ports are available on demand across the subnet.", font=font(22, True), fill=paint(NAVY), anchor="mm")
    allocations = ("16K ports", "16K ports", "16K ports")
    for index, y in enumerate((300, 420, 540)):
        draw.rounded_rectangle((1000, y, 1195, y + 70), radius=12, fill=paint(WHITE), outline=paint(BLUE), width=3)
        draw.text((1097, y + 35), f"VM {index + 1}", font=font(22, True), fill=paint(NAVY), anchor="mm")
        draw.rounded_rectangle((1280, y, 1635, y + 70), radius=12, fill=paint(WHITE), outline=paint(BLUE), width=3)
        draw.text((1457, y + 35), allocations[index], font=font(22, True), fill=paint(BLUE), anchor="mm")
        draw.line((1195, y + 35, 1280, y + 35), fill=paint(BLUE), width=5)
    draw.text((1332, 665), "A heavy talker can exhaust its allocation while peers remain idle.", font=font(22, True), fill=paint(NAVY), anchor="mm")
    image.save(path, quality=95)
    return path


def create_scenarios_diagram() -> Path:
    nodes = [
        ("Small admin VM", "Standard public IP\nDirect, exceptional, visible", "D83B01"),
        ("Autoscaling app", "NAT Gateway\nSubnet-wide predictable egress", TEAL),
        ("Regulated estate", "Private Link + Firewall\nPrivate PaaS and inspected internet", "C239B3"),
    ]
    return create_horizontal_flow(
        "worked-architectures.png",
        "The requirement, not fashion, selects the architecture",
        nodes,
        "Use the smallest architecture that meets security, scale, resilience, and observability requirements.",
    )


def create_discovery_diagram() -> Path:
    nodes = [
        ("Inventory", "Subnets, NICs, routes,\npublic IPs, NAT, LB, Firewall", BLUE),
        ("Observe", "VNet flow logs, Traffic Analytics,\nmetrics and application telemetry", TEAL),
        ("Classify", "Required, private-capable,\nunknown, or unnecessary", "C239B3"),
        ("Decide", "Choose explicit path and\ncapacity from evidence", "107C10"),
    ]
    return create_horizontal_flow(
        "discovery-evidence-flow.png",
        "Discovery turns an implicit dependency into evidence",
        nodes,
        "A resource inventory shows possible paths; traffic observation shows which paths the workload actually needs.",
    )


def create_migration_diagram() -> Path:
    nodes = [
        ("1. Observe", "Representative traffic window", BLUE),
        ("2. Build egress", "Deploy and capacity-plan", TEAL),
        ("3. Validate", "Test destinations and egress IP", "107C10"),
        ("4. Disable", "defaultOutboundAccess = false", "C239B3"),
        ("5. Deallocate", "Stop/deallocate, start, verify", "D83B01"),
    ]
    return create_horizontal_flow(
        "migration-sequence.png",
        "The order is the safety mechanism",
        nodes,
        "Explicit egress first. Disable the fallback second. Deallocate existing VMs third.",
    )


def create_decision_diagram() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "outbound-decision-tree.png"
    image = Image.new("RGB", (1800, 1060), paint(WHITE))
    draw = ImageDraw.Draw(image)
    draw.text((900, 52), "A practical outbound architecture decision tree", font=font(40, True), fill=paint(NAVY), anchor="mm")
    boxes = [
        ((600, 130, 1200, 275), "Does the workload need a public endpoint?", BLUE),
        ((80, 375, 650, 535), "No, the service supports Private Link", "107C10"),
        ((1150, 375, 1720, 535), "Yes, public access is required", "D83B01"),
        ((900, 650, 1460, 810), "Need central inspection or FQDN policy?", "C239B3"),
        ((190, 875, 720, 1020), "NAT Gateway\nGeneral-purpose subnet egress", TEAL),
        ((1080, 875, 1610, 1020), "Firewall / NVA\nGoverned inspected egress", "C239B3"),
    ]
    for bounds, text, color in boxes:
        draw.rounded_rectangle(bounds, radius=18, fill=paint("F7F8FA"), outline=paint(color), width=4)
        draw.multiline_text(((bounds[0] + bounds[2]) / 2, (bounds[1] + bounds[3]) / 2), text, font=font(26, True), fill=paint(NAVY), anchor="mm", align="center")
    draw.line((600, 202, 365, 375), fill=paint("107C10"), width=6)
    draw.text((430, 295), "No", font=font(22, True), fill=paint("107C10"), anchor="mm")
    draw.line((1200, 202, 1435, 375), fill=paint("D83B01"), width=6)
    draw.text((1370, 295), "Yes", font=font(22, True), fill=paint("D83B01"), anchor="mm")
    draw.line((1435, 535, 1180, 650), fill=paint("C239B3"), width=6)
    draw.line((900, 730, 455, 875), fill=paint(TEAL), width=6)
    draw.text((690, 810), "No", font=font(22, True), fill=paint(TEAL), anchor="mm")
    draw.line((1460, 730, 1345, 875), fill=paint("C239B3"), width=6)
    draw.text((1450, 810), "Yes", font=font(22, True), fill=paint("C239B3"), anchor="mm")
    draw.multiline_text((365, 575), "Use private DNS and disable public access\nwhere the service supports it.", font=font(21), fill=paint(GRAY), anchor="mm", align="center", spacing=6)
    draw.multiline_text((1435, 575), "Load Balancer rules fit an existing backend pool;\na VM public IP is a narrow exception.", font=font(20), fill=paint(GRAY), anchor="mm", align="center", spacing=6)
    image.save(path, quality=95)
    return path


def add_picture(document: Document, path: Path, caption: str, alt_text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(6.25))
    drawing_properties = inline_shape._inline.docPr
    drawing_properties.set("descr", alt_text)
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def add_opening(document: Document) -> None:
    add_heading(document, "Start with the packet, not the product")
    document.add_paragraph(
        "A virtual machine at 10.20.1.4 needs to call a public API over HTTPS. "
        "The operating system creates a TCP connection from an ephemeral source port, such as "
        "10.20.1.4:51542, to the API at 198.51.100.25:443. The destination is public, but the "
        "source address is private and cannot be routed across the public internet. Something at "
        "the edge must translate that source identity."
    )
    add_picture(
        document,
        create_snat_diagram(),
        "Figure 1. SNAT replaces the private source tuple for the outbound path and preserves state for the reply.",
        "A workload VM sends a private packet to an explicit egress service, which translates the source before the packet reaches a public service. Return traffic follows stored flow state.",
    )


def add_default_outbound_section(document: Document) -> None:
    add_heading(document, "How Azure made outbound connectivity feel automatic")
    document.add_paragraph(
        "Historically, a VM deployed into a nonprivate Azure subnet could reach public endpoints even when its NIC had no public IP, its subnet had no NAT Gateway, and no Load Balancer outbound rule or firewall route existed. Azure supplied default outbound access: an implicit SNAT path through a Microsoft-owned public IP address."
    )
    add_picture(
        document,
        create_horizontal_flow(
            "legacy-default-outbound.png",
            "The legacy convenience hid an architectural dependency",
            [
                ("VM", "10.20.1.4\nNo public IP on NIC", BLUE),
                ("Azure platform", "Implicit SNAT\nNo customer resource", "D83B01"),
                ("Hidden public IP", "Microsoft-owned\nCan change without notice", "C239B3"),
                ("Internet", "Public API\nsees hidden source", GRAY),
            ],
            "The workload works, but the egress identity, capacity, and lifecycle are not explicitly owned by the service team.",
        ),
        "Figure 2. Default outbound access supplied an implicit public identity when no explicit outbound method existed.",
        "A VM reaches the internet through implicit platform SNAT and a hidden Microsoft-owned public IP.",
    )
    document.add_paragraph(
        "This did not assign a visible public IP resource to the NIC, and it did not permit arbitrary unsolicited inbound connections. The risk is not that outbound SNAT behaves like an open inbound listener. The risk is that internet reachability exists by default through a path the workload owner did not deliberately create, govern, capacity-plan, or monitor as a resource."
    )
    add_labeled_paragraph(document, "Assigned is not the same as used", "A default outbound IP can remain associated with a VM in a nonprivate subnet even after NAT Gateway or a firewall route is added. New flows can use the explicit path while the old default identity remains assigned as a fallback. Making the subnet private and stop/deallocating the VM removes that fallback state.")
    add_labeled_paragraph(document, "The 2026 behavior change", "For API behavior released after March 31, 2026, subnets in new VNets default to private, so an explicit method is required to reach public internet and Microsoft endpoints. Existing VNets were not automatically changed. Older API versions can still leave defaultOutboundAccess null, which implicitly permits default outbound access.")
    add_callout(document, "Precise language", "A private subnet is a subnet with defaultOutboundAccess explicitly set to false. It can still reach public endpoints through NAT Gateway, Load Balancer outbound rules, a VM public IP, or a routed firewall/NVA. Private means no platform-created fallback, not no outbound traffic.")


def add_why_section(document: Document) -> None:
    add_heading(document, "Why implicit egress is the wrong default")
    add_table(
        document,
        ["Concern", "What implicit access makes difficult", "What explicit egress provides"],
        [
            ["Zero Trust", "Reachability appears without an intentional architecture decision.", "A named path that can be reviewed, restricted, and removed."],
            ["Ownership", "The public IP is Microsoft-owned and can change without notice.", "Customer-managed IPs or prefixes with a deliberate lifecycle."],
            ["Capacity", "SNAT behavior is not sized to workload concurrency.", "Measurable port pools, scale limits, and capacity controls."],
            ["Operations", "Incidents depend on a hidden platform fallback.", "Metrics, logs, effective routes, and a documented next hop."],
            ["Partner allow-lists", "The observed source address is not a stable contract.", "Predictable, allow-listable egress addresses."],
        ],
        [1.0, 2.5, 2.5],
    )
    document.add_paragraph(
        "Default outbound access also has technical limitations: the default IP does not support fragmented packets or ICMP ping; multiple NICs can produce inconsistent source addresses; and scale-set changes can change observed egress identities. These are symptoms of the deeper issue: a production network contract should not depend on an implicit convenience path."
    )
    add_callout(document, "Do not confuse two controls", "An NSG answers whether a packet is allowed. An outbound architecture answers where an allowed packet goes and which source identity it uses. Secure egress normally needs both policy and a deliberate path.")


def add_architectures_section(document: Document) -> None:
    add_heading(document, "The explicit outbound architectures")
    document.add_paragraph(
        "Azure offers several explicit methods because workloads have different boundaries. The correct question is not which product is universally best. It is which scope should own egress, whether traffic requires inspection, how many concurrent flows must be translated, which IPs partners will allow, and how the design behaves during failure."
    )
    add_picture(
        document,
        create_options_diagram(),
        "Figure 3. Explicit outbound can be subnet-wide, backend-pool-scoped, instance-scoped, centrally routed, or avoided for supported PaaS traffic.",
        "A private workload connects to NAT Gateway, Load Balancer outbound rules, Firewall or NVA, an instance public IP, or a Private Link destination.",
    )

    add_heading(document, "1. NAT Gateway: the general-purpose subnet egress", level=2)
    document.add_paragraph(
        "Azure NAT Gateway attaches to one or more subnets in a single VNet. For new outbound flows from those subnets, it translates private workload addresses to customer-selected static public IP addresses or prefixes. Workloads do not need public IPs on their NICs, and unsolicited inbound connections cannot be initiated through the gateway."
    )
    add_bullets(document, [
        "Strengths: simple subnet attachment, predictable egress identity, dynamic SNAT-port allocation, high scale, and no route-table change for the normal internet system route.",
        "Tradeoffs: hourly and data-processing cost, one-VNet scope, SKU and regional constraints, and no destination filtering by itself. Use NSGs or a firewall when policy requires more than an egress identity.",
        "Precedence: a UDR to a virtual appliance or virtual network gateway overrides NAT Gateway. Otherwise NAT Gateway takes precedence over instance public IPs, Load Balancer outbound rules, and default access for new connections.",
        "Best fit: private application tiers, VM scale sets, AKS or App Service integrations where supported, and most subnets that need scalable public egress without centralized inspection.",
    ])
    add_table(
        document,
        ["NAT Gateway SKU", "Standard", "StandardV2"],
        [
            ["Availability zones", "Zonal or no-zone placement; not zone redundant", "Zone redundant by default"],
            ["IP support", "IPv4", "IPv4 and IPv6, plus NAT64"],
            ["Protocols", "TCP and UDP", "TCP, UDP, and ICMP echo"],
            ["Throughput", "Up to 50 Gbps", "Up to 100 Gbps"],
            ["Flow logs", "Use surrounding network telemetry", "NAT Gateway flow logs supported"],
            ["Migration caveat", "Uses Standard public IP resources", "Requires StandardV2 public IPs; no in-place Standard upgrade"],
        ],
        [1.4, 2.2, 2.4],
    )
    add_callout(document, "2026 design note", "StandardV2 changes the old NAT Gateway availability-zone and IPv6 comparison. It is zone redundant and dual stack, but it is not available in every region and has public-IP compatibility and migration constraints. Check current regional support before standardizing.")

    add_heading(document, "2. Instance-level public IP: direct and deliberately exceptional", level=2)
    document.add_paragraph(
        "A Standard public IP attached to a VM NIC creates a 1:1 public identity. Azure uses stateless 1:1 NAT rather than sharing a SNAT pool, so the VM has its ephemeral ports available. This is explicit and easy to troubleshoot, but it places public addressing at instance scope."
    )
    add_bullets(document, [
        "Strengths: direct identity, simple packet path, full per-instance ephemeral-port availability, and support for protocols that shared outbound NAT might not handle.",
        "Tradeoffs: one public resource per instance, larger governance surface, exposure that must be constrained with NSGs, and poor fit for autoscaling fleets or centralized allow-list management.",
        "Best fit: a justified administrative appliance, protocol-specific exception, or small standalone VM where instance identity is the requirement rather than a shortcut.",
    ])

    add_heading(document, "3. Load Balancer outbound rules: explicit egress for a backend pool", level=2)
    document.add_paragraph(
        "A Standard public Load Balancer can use its frontend public IPs for outbound SNAT from VMs in a backend pool. Outbound rules declare the protocol, idle timeout, TCP reset behavior, and SNAT-port allocation. This is useful when a load balancer already owns the backend boundary, but capacity must be planned against pool size and frontend IP count."
    )
    add_bullets(document, [
        "Strengths: reuses a load-balancer frontend, gives declarative control, can share a public IP prefix, and supports IPv6 outbound rules.",
        "Tradeoffs: ports are allocated to backend instances rather than dynamically shared across the subnet; a heavy talker can exhaust its allocation while another VM has unused ports.",
        "Constraints: outbound NAT supports TCP and UDP, primary IPv4 NIC configurations, and requires all relevant availability-set or scale-set instances in the backend pool. Idle timeout is configurable from 4 to 120 minutes.",
        "Important caveat: an IP-address-based backend pool can behave like Basic Load Balancer with default outbound enabled. Configure backend membership by NIC or use NAT Gateway for a secure-by-default, demanding workload.",
        "Best fit: an existing, stable Standard Load Balancer backend pool with understood concurrency and a reason to reuse its frontend identity.",
    ])
    add_picture(
        document,
        create_snat_capacity_diagram(),
        "Figure 4. NAT Gateway dynamically shares ports at subnet scope; Load Balancer outbound rules allocate capacity to backend instances.",
        "NAT Gateway workloads draw from a shared dynamic SNAT pool while Load Balancer backend VMs receive separate fixed allocations.",
    )

    add_heading(document, "4. Azure Firewall or NVA: governed egress through routing", level=2)
    document.add_paragraph(
        "When the requirement is not merely to translate traffic but to inspect and govern it, associate a route table with workload subnets and send 0.0.0.0/0 to Azure Firewall or another network virtual appliance. The appliance becomes the policy enforcement point and normally performs SNAT before traffic reaches the internet."
    )
    add_bullets(document, [
        "Strengths: centralized network and application rules, FQDN-aware policy where supported, threat intelligence, detailed logs, and a consistent hub-and-spoke control point.",
        "Tradeoffs: higher cost, added latency, route and DNS complexity, high-availability design, asymmetric-routing risk, and explicit SNAT-capacity planning.",
        "Scale option: Azure Firewall public IPs add SNAT capacity; NAT Gateway can be attached to AzureFirewallSubnet for a larger dynamically allocated pool. StandardV2 is the zone-redundant NAT option for a zone-redundant firewall design.",
        "Forced tunneling: internet traffic can be sent onward to an on-premises edge or another appliance. This changes SNAT and return-path behavior and must be designed as an end-to-end route, not just a spoke UDR.",
        "Best fit: regulated environments, destination allow-listing, centralized inspection, and organizations already operating a resilient hub network.",
    ])

    add_heading(document, "5. Private Link: remove public egress from the path", level=2)
    document.add_paragraph(
        "If the destination is Azure Storage, SQL, Key Vault, Container Registry, or another supported service, a Private Endpoint gives that service a private IP in the VNet. Traffic stays on private addressing and does not consume internet SNAT ports. This is often the strongest answer because the public path disappears rather than being filtered."
    )
    add_bullets(document, [
        "Strengths: private service access, no internet SNAT for that flow, service-level approval, and the ability to disable the service's public endpoint.",
        "Tradeoffs: per-service and per-subresource endpoints, private DNS design, endpoint lifecycle and cost, and service-specific limitations.",
        "Best fit: supported Azure PaaS dependencies. It complements NAT Gateway or Firewall for package feeds, external APIs, and other destinations that still require a public endpoint.",
    ])


def add_scenarios_section(document: Document) -> None:
    add_heading(document, "Three architectures, three different reasons")
    add_picture(
        document,
        create_scenarios_diagram(),
        "Figure 5. A small exception, a scalable application tier, and a regulated estate should not be forced into the same egress pattern.",
        "Three scenarios map a small administrative VM to a public IP, an autoscaling application to NAT Gateway, and a regulated estate to Private Link plus Firewall.",
    )
    add_heading(document, "Scenario A: the isolated administrative VM", level=2)
    add_labeled_paragraph(document, "Requirement", "One fixed VM must reach and be reached through a stable public identity for a protocol-specific administrative workflow. Scaling and centralized inspection are not requirements.")
    add_labeled_paragraph(document, "Choice", "A Standard public IP on the NIC, protected by a restrictive NSG and host firewall.")
    add_labeled_paragraph(document, "Why not NAT Gateway", "NAT Gateway would solve outbound identity but not the justified inbound identity. Adding it would create another resource without satisfying the complete requirement.")
    add_labeled_paragraph(document, "Failure to avoid", "Treating this exception as the fleet default. Every additional public NIC increases ownership, policy, and inventory work.")

    add_heading(document, "Scenario B: an autoscaling application subnet", level=2)
    add_labeled_paragraph(document, "Requirement", "Dozens of private instances call external payment and package endpoints. Instance count and connection concurrency vary, and partners allow-list a small public prefix.")
    add_labeled_paragraph(document, "Choice", "NAT Gateway on the workload subnet with a customer-owned public IP prefix. Private Endpoints handle Storage and Key Vault.")
    add_labeled_paragraph(document, "Why not Load Balancer rules", "The application has uneven talkers and frequent scale-out. Dynamic subnet-level port allocation is a better fit than fixed backend-instance allocations.")
    add_labeled_paragraph(document, "Return path", "The partner replies to the NAT public IP and translated port; NAT state maps the response to the initiating private instance. No inbound connection can begin through the gateway.")

    add_heading(document, "Scenario C: a regulated hub-and-spoke estate", level=2)
    add_labeled_paragraph(document, "Requirement", "All public destinations require application-aware policy, central logging, incident correlation, and controlled change. Azure PaaS traffic should avoid the internet.")
    add_labeled_paragraph(document, "Choice", "Spoke UDRs send 0.0.0.0/0 to a zone-redundant Azure Firewall. Private Endpoints serve supported PaaS dependencies. NAT Gateway StandardV2 on AzureFirewallSubnet is considered if Firewall SNAT scale requires it.")
    add_labeled_paragraph(document, "Why not NAT Gateway alone", "NAT Gateway supplies identity and translation, not FQDN filtering or centralized application policy.")
    add_labeled_paragraph(document, "Failure to avoid", "Routing only the forward path. Hub, peering, gateway propagation, appliance SNAT, and return routes must remain symmetric.")


def add_decision_section(document: Document) -> None:
    add_heading(document, "Choose from requirements, not from habit")
    add_picture(
        document,
        create_decision_diagram(),
        "Figure 6. Start by removing unnecessary public paths, then decide whether remaining internet traffic requires centralized inspection.",
        "A decision tree routes private-capable services to Private Link, general public egress to NAT Gateway, and inspected egress to Firewall or NVA.",
    )
    add_table(
        document,
        ["Method", "Scope", "SNAT model", "Inspection", "Best fit"],
        [
            ["NAT Gateway", "Subnet", "Dynamic shared pool", "No destination policy by itself", "Scalable general-purpose egress"],
            ["VM public IP", "NIC/IP configuration", "1:1 NAT", "NSG/host controls", "Small justified instance exception"],
            ["LB outbound rules", "Backend pool", "Per-instance allocation", "NSG; no application inspection", "Existing stable public LB pool"],
            ["Firewall/NVA", "Routed subnets/spokes", "Appliance-managed", "Central network/application policy", "Governed or regulated egress"],
            ["Private Link", "Service endpoint", "No internet SNAT", "Private service boundary", "Supported Azure PaaS"],
        ],
        [1.1, 1.1, 1.3, 1.6, 1.7],
    )


def add_discovery_section(document: Document) -> None:
    add_heading(document, "Discover before you disable")
    document.add_paragraph(
        "A subnet property tells you whether the fallback is permitted. It does not tell you why an application needs egress. Migration therefore needs two forms of evidence: configuration inventory and observed traffic."
    )
    add_picture(
        document,
        create_discovery_diagram(),
        "Figure 7. Inventory identifies possible paths; observation reveals real dependencies; classification turns both into an architecture decision.",
        "A four-step flow moves from resource inventory to traffic observation, dependency classification, and explicit architecture selection.",
    )
    add_heading(document, "First, inventory the possible paths", level=2)
    add_bullets(document, [
        "Read defaultOutboundAccess on every subnet. Treat null as a legacy-compatible state, not as proof of private behavior.",
        "Map NICs, VMs, VM scale-set instances, and their IP configurations to subnets.",
        "Detect NIC public IPs, subnet NAT Gateway associations, Standard Load Balancer backend membership and outbound rules, UDRs, Azure Firewall or NVA next hops, service endpoints, and Private Endpoints.",
        "Use Azure Advisor recommendations to find VM and VMSS NICs flagged with default outbound connectivity enabled.",
    ])
    add_heading(document, "Then, observe what the workload actually does", level=2)
    document.add_paragraph(
        "Enable Network Watcher virtual network flow logs for a representative business window and send them to a standard Storage account. Traffic Analytics can enrich and aggregate the records in Log Analytics. Each Layer 4 flow includes source and destination IPs, ports, protocol, direction, state, packet and byte counts, and the rule or platform decision associated with the flow."
    )
    add_callout(document, "Use VNet flow logs for new work", "NSG flow logs can no longer be newly created and retire on September 30, 2027. Virtual network flow logs are the current direction, simplify VNet-wide coverage, and avoid enabling separate logging at every NSG attachment.")
    add_labeled_paragraph(document, "What flow logs do not know", "An IP and port do not explain business intent. CDN addresses change, several hostnames can share an IP, and encrypted HTTPS hides the URL. Correlate flow evidence with DNS logs where available, firewall application logs, dependency documentation, and application telemetry before creating an allow-list.")
    add_labeled_paragraph(document, "Classification outcome", "Mark each dependency as required public, replaceable with Private Link, internal through peering or private WAN, obsolete, or unknown. Unknown is a reason to investigate, not a reason to reproduce unrestricted internet access forever.")


def add_migration_section(document: Document) -> None:
    add_heading(document, "A migration sequence designed not to surprise the application")
    add_picture(
        document,
        create_migration_diagram(),
        "Figure 8. The explicit path is deployed and tested before the subnet fallback is disabled; existing VMs are deallocated only after that validation.",
        "A five-step sequence observes traffic, builds explicit egress, validates it, disables default outbound access, then stop-deallocates and restarts existing VMs.",
    )
    add_table(
        document,
        ["Stage", "Action", "Exit evidence"],
        [
            ["Observe", "Capture a representative period including patching, deployments, certificate checks, backups, and scale events.", "Destinations and peak connection behavior are understood."],
            ["Reduce", "Move supported PaaS dependencies to Private Endpoints and remove obsolete calls.", "Only justified public dependencies remain."],
            ["Build", "Deploy the selected public IPs/prefixes, NAT Gateway, LB rules, or routed Firewall/NVA path.", "Infrastructure exists and capacity assumptions are documented."],
            ["Validate", "Test DNS, required protocols, partner allow-lists, egress IP, routes, logging, metrics, and failure behavior.", "The explicit path works before the fallback changes."],
            ["Disable", "Set defaultOutboundAccess to false through governed infrastructure as code.", "The subnet is private and new fallback assignments are prevented."],
            ["Refresh", "Stop/deallocate and start existing VMs or roll scale-set instances in controlled waves.", "NIC defaultOutboundConnectivityEnabled state and Advisor findings clear."],
        ],
        [0.9, 3.4, 2.1],
    )
    document.add_paragraph(
        "The stop operation must deallocate the VM; a guest restart is not enough. Deallocation allows Azure to update the NIC-level default outbound connectivity state. Plan for workload availability, instance repair, health probes, and stateful services before scheduling the wave."
    )
    add_callout(document, "The dangerous ordering", "Setting defaultOutboundAccess to false before required egress exists can break Windows activation and updates, Linux package repositories, container registries, certificate revocation checks, telemetry agents, external APIs, and deployment tooling. Build and test the replacement first.")


def add_validation_section(document: Document) -> None:
    add_heading(document, "Validate the path from both ends")
    add_bullets(document, [
        "From the workload, resolve each required hostname and establish the real TCP or UDP connection. A generic web test does not prove that the application dependency works.",
        "From an external endpoint you control, record the observed source IP and confirm that it belongs to the intended public IP or prefix.",
        "Inspect effective routes and Network Watcher next-hop results. Confirm that 0.0.0.0/0 resolves to the intended internet, virtual appliance, or virtual network gateway path.",
        "Check NAT Gateway, Load Balancer, or Firewall health and SNAT metrics during a representative concurrency test. Connection reuse in the application remains important even with a scalable gateway.",
        "Confirm flow or firewall logs contain the expected source subnet, destination, protocol, decision, and bytes. Observability must survive the migration, not merely connectivity.",
        "After stop/deallocate and start, recheck Azure Advisor and NIC state. An explicit route carrying traffic does not by itself remove the legacy default assignment.",
    ])
    add_heading(document, "Common traps", level=2)
    add_table(
        document,
        ["Symptom", "Likely cause", "What to inspect"],
        [
            ["Old sessions use one IP; new sessions use another", "Long-lived flows preserve their original path while new flows take the new higher-precedence method.", "Connection age, flow state, and cutover timing."],
            ["Private subnet cannot reach a service-tag route with next hop Internet", "That route depended on default outbound access rather than an explicit translator.", "UDR next hop and an explicit outbound method for the source."],
            ["Firewall sees traffic but replies fail", "Asymmetric route, missing SNAT, or incorrect gateway propagation.", "Forward and return effective routes across spokes and hub."],
            ["New scale-set instances fail package download", "The explicit method is missing from the instance subnet or partner allow-list.", "Subnet association, observed egress IP, and destination policy."],
            ["NAT Gateway association is rejected", "The subnet contains incompatible Basic SKU resources or is a restricted managed/delegated subnet.", "SKU compatibility and service-specific networking constraints."],
        ],
        [1.6, 2.5, 2.0],
    )


def add_conclusion_and_sources(document: Document) -> None:
    add_heading(document, "What to remember")
    add_bullets(document, [
        "Outbound access is an architecture decision, not a convenience attached to a VM.",
        "A private subnet means no implicit platform egress. It can still have carefully designed public connectivity.",
        "Private Link removes public egress for supported Azure services; NAT Gateway or Firewall handles the public destinations that remain.",
        "NAT Gateway is the default general-purpose choice, Load Balancer rules fit a deliberate backend-pool boundary, a VM public IP is an exception, and Firewall/NVA is for inspection and governance.",
        "The safe migration order is explicit egress, validation, defaultOutboundAccess=false, then stop/deallocate and start existing VMs.",
    ])
    document.add_paragraph(
        "The real improvement is not simply changing a Boolean property. It is replacing an invisible assumption with a network contract the team can explain: which traffic may leave, through which next hop, under which identity and policy, with how much capacity, and how an operator proves that it is working."
    )
    add_heading(document, "Official references", level=1)
    document.add_paragraph("Product behavior and limits in this article were verified against Microsoft Learn on August 5, 2026.")
    sources = [
        ("Default outbound access in Azure", "https://learn.microsoft.com/azure/virtual-network/ip-services/default-outbound-access"),
        ("What is Azure NAT Gateway?", "https://learn.microsoft.com/azure/nat-gateway/nat-overview"),
        ("Azure NAT Gateway SKUs", "https://learn.microsoft.com/azure/nat-gateway/nat-sku"),
        ("Azure Load Balancer outbound connections and SNAT", "https://learn.microsoft.com/azure/load-balancer/load-balancer-outbound-connections"),
        ("Azure Load Balancer outbound rules", "https://learn.microsoft.com/azure/load-balancer/outbound-rules"),
        ("Virtual network flow logs", "https://learn.microsoft.com/azure/network-watcher/vnet-flow-logs-overview"),
        ("NSG flow logs retirement guidance", "https://learn.microsoft.com/azure/network-watcher/nsg-flow-logs-overview"),
        ("Azure Firewall and NAT Gateway integration", "https://learn.microsoft.com/azure/firewall/integrate-with-nat-gateway"),
        ("Azure Firewall forced tunneling", "https://learn.microsoft.com/azure/firewall/forced-tunneling"),
        ("Azure Private Endpoint overview", "https://learn.microsoft.com/azure/private-link/private-endpoint-overview"),
    ]
    for title, url in sources:
        paragraph = document.add_paragraph()
        paragraph.add_run("• ")
        add_hyperlink(paragraph, title, url)
    document.add_paragraph(
        "Source Network Address Translation (SNAT) replaces the packet's private source IP and "
        "port with a public IP and an available SNAT port. The translator records the mapping. "
        "When the response arrives, that state tells the platform which private workload should "
        "receive it. Unsolicited inbound traffic has no matching state and is not admitted by this "
        "outbound-only translation."
    )
    add_callout(
        document,
        "Core idea",
        "Outbound connectivity is a designed packet path with an owned public identity, routing, "
        "translation capacity, policy, and observability. It should not be an undocumented side "
        "effect of placing a VM in a virtual network.",
    )


def add_introductory_article(document: Document) -> None:
    add_heading(document, "What outbound connectivity actually means")
    document.add_paragraph(
        "Imagine a private virtual machine that needs to download an operating-system update or "
        "call a payment API. The VM starts the conversation, so the traffic is outbound. Its source "
        "address might be 10.20.1.4, an address that works inside the virtual network but cannot be "
        "used on the public internet. Before the packet leaves Azure, that private source must be "
        "represented by a public address."
    )
    add_picture(
        document,
        create_snat_diagram(),
        "Figure 1. An outbound service gives a private workload a public identity for the connection and remembers where the reply belongs.",
        "A private virtual machine sends a packet through an outbound service, which translates its source before reaching a public service.",
    )
    document.add_paragraph(
        "This translation is called Source Network Address Translation, or SNAT. The outbound service "
        "temporarily maps the VM's private address and port to a public address and port. When the "
        "reply returns, the saved mapping sends it back to the correct VM. This does not mean the "
        "internet can freely start connections to the VM; the return packet is accepted because the "
        "VM started the conversation."
    )
    add_callout(
        document,
        "Simple mental model",
        "Outbound connectivity is the controlled exit door for a private workload. Security depends "
        "on knowing where that door is, who owns it, and what traffic is allowed through it.",
    )

    add_heading(document, "The old convenience: internet access without an explicit door")
    document.add_paragraph(
        "For many years, Azure made this feel automatic. If a VM had no public IP, NAT Gateway, Load "
        "Balancer outbound rule, or firewall path, Azure could still give it default outbound access "
        "through a Microsoft-owned public IP. The VM reached the internet even though the team had not "
        "created an outbound resource."
    )
    add_picture(
        document,
        create_horizontal_flow(
            "legacy-default-outbound.png",
            "Default outbound access: useful, but implicit",
            [
                ("Private VM", "Starts an internet connection", BLUE),
                ("Azure platform", "Supplies hidden SNAT", "D83B01"),
                ("Public service", "Sees a Microsoft-owned IP", GRAY),
            ],
            "The application works, but the team did not deliberately create or own its internet exit.",
        ),
        "Figure 2. Default outbound access hid the public egress path behind platform behavior.",
        "A private virtual machine reaches a public service through hidden Azure platform translation.",
    )
    document.add_paragraph(
        "The problem is not that every VM suddenly becomes an inbound web server. The security problem "
        "is that public reachability exists by default. A compromised process can call an unknown "
        "internet destination through the same invisible path, and operators may not have a stable "
        "public IP, a central inspection point, or even a clear diagram showing how the traffic leaves."
    )
    add_table(
        document,
        ["Implicit default", "Explicit outbound"],
        [
            ["Internet access appears automatically", "Internet access is an intentional design decision"],
            ["Public identity is Microsoft-owned and can change", "The team owns or selects the public identity"],
            ["The path is easy to overlook", "The path can be reviewed, logged, and governed"],
            ["A fallback can survive unnoticed", "Removing the outbound resource removes the path"],
        ],
        [3.0, 3.0],
    )
    add_callout(
        document,
        "Private does not mean isolated",
        "Setting defaultOutboundAccess to false removes Azure's implicit fallback. A private subnet "
        "can still reach approved public destinations through an explicit NAT Gateway or firewall.",
    )

    add_heading(document, "Why disabling the default improves security")
    document.add_paragraph(
        "Zero Trust starts with an uncomplicated question: why should this workload be allowed to "
        "communicate? Default outbound access answers that question with 'because nothing explicitly "
        "stopped it.' Explicit outbound connectivity changes the answer to 'because we created a "
        "known path for a known need.'"
    )
    add_bullets(document, [
        "Intent becomes visible. A network diagram and deployment template show that internet access is required.",
        "The public identity becomes predictable. Partners can allow-list an address owned by the workload team.",
        "Monitoring has a natural place. NAT or firewall logs can show which destinations workloads contact.",
        "Policy can be stronger. A firewall can restrict destinations instead of giving every allowed packet the same unrestricted exit.",
        "Removing access becomes deliberate. A subnet without an explicit outbound method does not silently retain an internet fallback.",
    ])
    document.add_paragraph(
        "As of March 31, 2026, new Azure virtual networks use private subnets by default under the new "
        "API behavior. Existing virtual networks are not automatically changed, which is why older "
        "environments still need to identify and replace their implicit dependency."
    )

    add_heading(document, "The main choices, without the product manual")
    add_picture(
        document,
        create_options_diagram(),
        "Figure 3. Different outbound methods solve different security and ownership problems.",
        "A private workload can use NAT Gateway, Load Balancer outbound rules, Firewall, a public IP, or Private Link depending on its need.",
    )
    add_table(
        document,
        ["Choice", "Think of it as", "Good introductory use case", "Security tradeoff"],
        [
            ["NAT Gateway", "A shared exit for a subnet", "Private application servers need predictable internet access", "Stable identity, but it does not decide which websites are allowed"],
            ["Azure Firewall", "A guarded and inspected exit", "A regulated service must restrict and log destinations", "More control, with more cost and operational complexity"],
            ["VM public IP", "A personal street address", "A small, justified standalone system needs its own identity", "Simple, but increases per-VM public exposure and governance"],
            ["Load Balancer rules", "An exit shared by one backend pool", "An existing load-balanced application reuses its frontend IP", "Explicit, but less flexible for uneven or rapidly scaling workloads"],
            ["Private Link", "No public exit for that service", "A VM reaches Storage, SQL, or Key Vault privately", "Strongest public-path reduction, but only for supported services"],
        ],
        [1.1, 1.3, 1.9, 2.0],
    )
    document.add_paragraph(
        "For a typical private application subnet that simply needs a stable way to call public "
        "services, NAT Gateway is usually the easiest explicit choice. If the organization must "
        "control destinations or inspect traffic, Azure Firewall is the more appropriate boundary. "
        "For supported Azure services, Private Link can remove the public journey entirely."
    )

    add_heading(document, "Two examples that make the difference concrete")
    add_heading(document, "Example 1: an application server downloading updates", level=2)
    document.add_paragraph(
        "A team runs ten private VMs. They need operating-system updates and must call one external API. "
        "With default outbound access, the calls work through a hidden public identity. The team instead "
        "attaches a NAT Gateway to the subnet and uses its known public IP. The VMs remain private, the "
        "external API can allow-list one stable address, and the internet path now exists because the "
        "team deliberately created it."
    )
    add_labeled_paragraph(
        document,
        "Security lesson",
        "NAT Gateway makes the exit explicit and predictable, but an NSG or NAT Gateway alone does not "
        "provide website-level filtering. If the VMs must reach only approved destinations, place the "
        "policy at a firewall or another suitable control point.",
    )

    add_heading(document, "Example 2: a compromised VM trying to send data out", level=2)
    document.add_paragraph(
        "Suppose an attacker gains control of an application process and tries to send data to an "
        "unfamiliar internet domain. Default outbound access may provide a route even though nobody "
        "designed one. A NAT Gateway makes the source identity visible, but it still permits general "
        "outbound connections unless other controls restrict them. A centrally governed firewall can "
        "allow required destinations, deny the unfamiliar domain, and record the attempt for investigation."
    )
    add_labeled_paragraph(
        document,
        "Security lesson",
        "Disabling default access removes an accidental path. It does not automatically create least "
        "privilege. The explicit replacement must still match the security requirement.",
    )

    add_picture(
        document,
        create_decision_diagram(),
        "Figure 4. First avoid public access where possible; then decide whether remaining internet traffic needs inspection.",
        "A simple decision tree recommends Private Link for supported private services, NAT Gateway for general egress, and Firewall for inspected egress.",
    )

    add_heading(document, "A safe way to think about the change")
    document.add_paragraph(
        "Before disabling the default, ask what the workload genuinely needs to reach. Move supported "
        "Azure services to Private Link where that makes sense. Give the remaining public traffic an "
        "explicit path, test it, and only then set defaultOutboundAccess to false. Existing VMs may need "
        "to be stopped and deallocated before their old default outbound identity is fully removed."
    )
    add_callout(
        document,
        "Remember the order",
        "Understand the need → create the explicit path → test it → disable the default. Disabling first "
        "can break updates, package downloads, monitoring agents, and external APIs.",
    )

    add_heading(document, "The idea to take away")
    document.add_paragraph(
        "The goal is not to make every private workload unable to use the internet. The goal is to stop "
        "internet access from appearing as an invisible default. Once the exit is explicit, teams can "
        "choose a stable identity, add the right inspection and logging, and explain why the connection "
        "is allowed. That is the security improvement: replacing convenience by accident with access by design."
    )

    add_heading(document, "Further reading")
    document.add_paragraph("These Microsoft Learn pages cover the implementation details when you are ready to apply the design:")
    sources = [
        ("Default outbound access in Azure", "https://learn.microsoft.com/azure/virtual-network/ip-services/default-outbound-access"),
        ("What is Azure NAT Gateway?", "https://learn.microsoft.com/azure/nat-gateway/nat-overview"),
        ("Azure Firewall overview", "https://learn.microsoft.com/azure/firewall/overview"),
        ("Azure Private Endpoint overview", "https://learn.microsoft.com/azure/private-link/private-endpoint-overview"),
    ]
    for title, url in sources:
        paragraph = document.add_paragraph()
        paragraph.add_run("• ")
        add_hyperlink(paragraph, title, url)


def configure_core_properties(document: Document) -> None:
    properties = document.core_properties
    properties.title = "Disabling Default Outbound Connectivity in Azure"
    properties.subject = "A visual guide to explicit outbound architectures and safe migration"
    properties.author = "Nikhil Kumar"
    properties.keywords = "Azure, outbound connectivity, SNAT, NAT Gateway, Azure Firewall, Zero Trust"
    properties.comments = "Technical guidance verified against Microsoft Learn in August 2026."


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("273444")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for level, size in ((1, 18), (2, 14), (3, 12)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Segoe UI"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(16 if level == 1 else 11)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Nikhil Kumar  |  Technical Blog  |  August 2026")
    run.font.name = "Segoe UI"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def build() -> Path:
    document = Document()
    configure_document(document)
    configure_core_properties(document)
    add_title(
        document,
        "Disabling Default Outbound Connectivity in Azure",
        "An introductory security guide to replacing invisible internet access with an explicit, controlled path",
    )
    add_callout(
        document,
        "The short version",
        "Private subnet does not mean no internet. It means no implicit internet path. If a workload "
        "needs a public endpoint, create and validate an explicit egress path before setting "
        "defaultOutboundAccess to false.",
    )
    add_introductory_article(document)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())